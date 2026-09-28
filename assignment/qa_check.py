# -*- coding: utf-8 -*-
"""QA：檢查 PPT 圖形是否越界、文字是否可能溢出；檢查 Word 結構順序。"""
import os
from pptx import Presentation
from pptx.util import Emu
from docx import Document

OUT = os.path.dirname(os.path.abspath(__file__))
EMU_IN = 914400


def check_pptx(path):
    prs = Presentation(path)
    W = prs.slide_width / EMU_IN
    H = prs.slide_height / EMU_IN
    print(f'PPT size: {W:.2f} x {H:.2f} in, slides={len(prs.slides)}')
    issues = 0
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left is None or sh.top is None:
                continue
            l, t = sh.left / EMU_IN, sh.top / EMU_IN
            w, h = (sh.width or 0) / EMU_IN, (sh.height or 0) / EMU_IN
            name = sh.shape_type
            if l < -0.02 or t < -0.02 or l + w > W + 0.02 or t + h > H + 0.02:
                print(f'  [OUT OF BOUNDS] slide {i}: {name} at ({l:.2f},{t:.2f}) '
                      f'size {w:.2f}x{h:.2f} -> right={l+w:.2f} bottom={t+h:.2f}')
                issues += 1
            # 粗略文字溢出估算
            if sh.has_text_frame:
                txt = sh.text_frame.text
                if not txt.strip():
                    continue
                sizes = [r.font.size.pt for p in sh.text_frame.paragraphs
                         for r in p.runs if r.font.size]
                fs = max(sizes) if sizes else 14
                chars_per_line = max(1, int(w * 96 / (fs * 0.52)))
                est_lines = 0
                for para in txt.split('\n'):
                    est_lines += max(1, -(-len(para) // chars_per_line))
                need_h = est_lines * fs * 1.32 / 72
                if need_h > h + 0.12:
                    print(f'  [MAY OVERFLOW] slide {i}: needs ~{need_h:.2f}in, box h={h:.2f}in '
                          f'(fs={fs}, chars/line={chars_per_line}) — "{txt[:48]}..."')
                    issues += 1
    print(f'PPT issues: {issues}')


def check_docx(path):
    doc = Document(path)
    print('\nDOCX structure:')
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        st = p.style.name
        if st.startswith('Heading'):
            print(f'  [{st}] {t[:70]}')
    pics = [s for s in doc.inline_shapes]
    print(f'  inline images: {len(pics)}')
    print(f'  total paragraphs: {len(doc.paragraphs)}')


if __name__ == '__main__':
    check_pptx(os.path.join(OUT, 'SAAS_Ideation_Presentation.pptx'))
    check_docx(os.path.join(OUT, 'SAAS_Individual_Ideation.docx'))
