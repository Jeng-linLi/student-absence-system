# -*- coding: utf-8 -*-
"""把 content.md 渲染成 Word 作业 (.docx)，并嵌入 Figure 1。"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.dirname(os.path.abspath(__file__))
MD  = os.path.join(OUT, 'content.md')
DOC = os.path.join(OUT, 'SAAS_Individual_Ideation.docx')
FIG = os.path.join(OUT, 'fig1_workflow.png')

INK = RGBColor(0x20, 0x1F, 0x1E)
MUTE = RGBColor(0x6B, 0x6B, 0x6B)
BRAND = RGBColor(0x0F, 0x6C, 0xBD)


def add_p(doc, text, *, size=11, bold=False, italic=False, color=INK,
          align=None, space_after=4, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_h1(doc, text):
    p = doc.add_heading(level=1)
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.color.rgb = INK; r.font.size = Pt(17); r.font.name = 'Calibri'
    return p


def add_h2(doc, text):
    p = doc.add_heading(level=2)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.color.rgb = INK; r.font.size = Pt(13.5); r.font.name = 'Calibri'
    return p


def add_inline(p, text, *, bold=False):
    r = p.add_run(text)
    r.font.size = Pt(11); r.font.name = 'Calibri'
    r.font.bold = bold; r.font.color.rgb = INK
    return r


def render_body(doc, text):
    """解析 Markdown 段落（**bold** 與簡單換行）。"""
    for raw in text.strip().split('\n\n'):
        raw = raw.strip()
        if not raw: continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        # split into inline tokens on **bold**
        parts = re.split(r'(\*\*[^*]+\*\*)', raw)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                add_inline(p, part[2:-2], bold=True)
            else:
                add_inline(p, part)


def main():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    # title block
    add_p(doc, 'Individual Ideation Assignment', size=10, color=MUTE, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    p = doc.add_paragraph()
    r = p.add_run('SAAS: A Student Absence and Academic Accommodation System')
    r.font.size = Pt(18); r.font.bold = True; r.font.name = 'Calibri'; r.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(4)
    add_p(doc, 'Johnny, Jeng-lin Li · HKUST(GZ)',
          size=10.5, color=MUTE, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)

    # parse content
    src = open(MD, encoding='utf-8').read()
    body, refs = src.split('## References', 1)
    title_line, rest = body.split('# SAAS', 1)[1].split('\n', 1)
    # title already used; skip
    rest = rest.lstrip('\n')

    # split on "\n## " so each block keeps its leading "## <heading>"
    blocks = re.split(r'\n(?=## )', rest)
    section_buffers = []  # (heading, body)
    for blk in blocks:
        if not blk.strip(): continue
        line, _, body_txt = blk.partition('\n')
        line = line.lstrip('#').strip()
        if re.match(r'^\d+\.\s', line):
            section_buffers.append((line, body_txt.strip()))
        elif line == 'Final Reflection':
            section_buffers.append(('Final Reflection', body_txt.strip()))

    for heading, body_text in section_buffers:
        if heading == 'Final Reflection':
            add_h2(doc, 'Final Reflection')
            render_body(doc, body_text)
            continue
        add_h2(doc, heading)
        render_body(doc, body_text)
        if 'Proposed Innovation' in heading:
            fig_p = doc.add_paragraph()
            fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_p.add_run().add_picture(FIG, width=Inches(5.8))
            cap = doc.add_paragraph()
            cr = cap.add_run('Figure 1. End-to-end flow of a SAAS request: one submission resolves '
                             'affected assessments and routes by policy through Cases A, B or C '
                             'before an instructor chooses an accommodation.')
            cr.font.size = Pt(9.5); cr.font.italic = True; cr.font.color.rgb = MUTE; cr.font.name = 'Calibri'
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)

    # references
    add_h1(doc, 'References')
    refs_body = refs.strip()
    for raw in refs_body.split('\n'):
        if raw.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(raw[2:].strip())
            r.font.size = Pt(10.5); r.font.name = 'Calibri'; r.font.color.rgb = INK
            p.paragraph_format.space_after = Pt(3)

    # footer note
    add_p(doc, 'Word count of the main body (Sections 1–5 plus Final Reflection, '
               'excluding the title block, figure caption and references): ≈ 1,030 words.',
          size=9.5, italic=True, color=MUTE, space_after=0)

    doc.save(DOC)
    print('saved', DOC)


if __name__ == '__main__':
    main()