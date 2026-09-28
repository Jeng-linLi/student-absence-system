# -*- coding: utf-8 -*-
"""SAAS Individual Ideation · PPT 簡報"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(OUT, 'SAAS_Ideation_Presentation.pptx')
FIG1 = os.path.join(OUT, 'fig1_workflow.png')
FIG2 = os.path.join(OUT, 'fig2_scurve.png')
FIG3 = os.path.join(OUT, 'fig3_complements.png')

BRAND = RGBColor(0x0F, 0x6C, 0xBD)
INK   = RGBColor(0x20, 0x1F, 0x1E)
MUTE  = RGBColor(0x6B, 0x6B, 0x6B)
LINE  = RGBColor(0xE1, 0xDF, 0xDD)
GREEN = RGBColor(0x0E, 0x7A, 0x3D)
AMBER = RGBColor(0xB8, 0x86, 0x0B)
RED   = RGBColor(0xB1, 0x0E, 0x1C)
PURPLE= RGBColor(0x6C, 0x4A, 0xB6)


def add_textbox(slide, x, y, w, h, text, *, size=14, bold=False, color=INK,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font='Calibri'):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = font; r.font.color.rgb = color
    return tb


def add_bullets(slide, x, y, w, h, items, *, size=13, color=INK, leading=1.3):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = leading
        r = p.add_run(); r.text = item
        r.font.size = Pt(size); r.font.name = 'Calibri'; r.font.color.rgb = color
        p.space_after = Pt(6)
    return tb


def add_rect(slide, x, y, w, h, *, fill=BRAND, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    sh.shadow.inherit = False
    return sh


def add_title_bar(slide, title_text, subtitle_text=None, *, accent=BRAND):
    # top accent bar
    add_rect(slide, 0.45, 0.34, 0.18, 0.42, fill=accent)
    add_textbox(slide, 0.7, 0.3, 12, 0.5, title_text, size=22, bold=True, color=INK)
    if subtitle_text:
        add_textbox(slide, 0.7, 0.78, 12, 0.32, subtitle_text, size=11, color=MUTE)
    add_rect(slide, 0.45, 1.15, 12.55, 0.02, fill=LINE)


def footer(slide, num, total):
    add_textbox(slide, 0.45, 7.0, 8, 0.25,
                'SAAS · Individual Ideation · Johnny, Jeng-lin Li', size=9, color=MUTE)
    add_textbox(slide, 11.6, 7.0, 1.4, 0.25, f'{num} / {total}',
                size=9, color=MUTE, align=PP_ALIGN.RIGHT)


NOTES = {
    1: "Hi, I'm Johnny — Jeng-lin Li — from HKUST(GZ). This is my Individual Ideation assignment. I'll propose "
       "a system that replaces today's email-based short-term leave with one policy-driven platform for "
       "universities.",
    2: "Today, asking for two days off is an email exercise. Students email each instructor separately, "
       "attach the same medical certificate several times, and negotiate each assessment individually. "
       "The programme office sees nothing; when a dispute arises, the Registry has no record.",
    3: "SAAS replaces the email ritual with one submission. The technologies are all mature; the new thing "
       "is the information design. Unit of work is not the request but its consequence. Policy lives in "
       "Registry-owned configuration. And privacy is structural — instructors see the accommodation to apply, "
       "never the medical reason or certificate behind it.",
    4: "Walk through the diagram: student submits one form → the system resolves enrolments against the "
       "assessment calendar and creates one conflict record per affected assessment → the request routes by "
       "policy through Case A, B or C → instructor picks one of 5 accommodations per conflict → the student "
       "sees every outcome in one timeline. Every step is audited.",
    5: "Three routing cases, all driven by configuration rather than code. Case A: medical → affected "
       "instructors. Case B: official activity → roster auto-verified, instructors notified only — the "
       "highest priority lane and no document required. Case C: any absence over seven days must pass "
       "the programme office and the Registry before instructors act.",
    6: "Apply Teece's complements argument. SAAS value is almost entirely co-produced. Without the assessment "
       "calendar, conflict detection is guesswork. Without campus SSO, adoption collapses. Without email "
       "and Teams delivery, decisions stall. Without Registry-owned configuration, the system fossilises. "
       "The hard work is securing the complements, not writing the application.",
    7: "Apply the S-curve. Rule-based workflow automation sits on the flat, mature part of its curve — "
       "competing on form features yields diminishing returns. SAAS matters only by starting a new curve. "
       "The same logic explains why MVP deliberately excludes AI triage: it sits on the steep part of the "
       "LLM curve, where reliability is still improving, and a wrong automated decision in a high-stakes "
       "academic process is unacceptable.",
    8: "The customer and payer is the university, not the student — Registry and IT buying an annual "
       "licence sized by student population. The value is that it makes an unmeasured cost measurable: "
       "instructor hours, dispute effort, compliance risk. Differentiator is the policy-configuration "
       "layer plus portability — one product, multiple institutions, different thresholds.",
    9: "The main positive impact is equity. Today the outcome depends on which instructor you ask and how "
       "persuasive your email is; a policy-driven process replaces discretion with consistent rules. "
       "But gains are uneven — need an assisted non-digital route. And a complete absence record is also "
       "a surveillance instrument. Mitigations are built in: instructors never see medical reasons; "
       "every decision stays human; AI is advisory, logged and refusable.",
    10: "Yes, develop further as a team project, but narrowed. The problem is real and bounded, the "
        "complements already exist on campus, and its hardest questions are organisational rather than "
        "technical — a genuine test of collaborative innovation. I'd scope the team project to the "
        "assessment-conflict workflow, where both the value and the privacy tension are sharpest.",
}


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]

    # ----------------- 1 Title
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, 13.333, 7.5, fill=RGBColor(0xFA, 0xFA, 0xFA))
    add_rect(s, 0, 6.6, 13.333, 0.9, fill=BRAND)
    add_rect(s, 0.55, 6.78, 0.16, 0.42, fill=RGBColor(0xFF, 0xFF, 0xFF))
    add_textbox(s, 0.45, 0.9, 12, 0.5, 'INDIVIDUAL IDEATION ASSIGNMENT',
                size=11, bold=True, color=MUTE)
    add_textbox(s, 0.45, 1.35, 12.4, 1.0,
                'SAAS — A Student Absence & Academic Accommodation System',
                size=30, bold=True, color=INK)
    add_textbox(s, 0.45, 2.3, 12, 0.5,
                'Replacing email-based absence with one policy-driven platform for universities',
                size=14, color=MUTE)
    # 4 mini callouts
    callouts = [
        ('9 user roles',    BRAND),
        ('9 leave categories', GREEN),
        ('5 instructor accommodations', PURPLE),
        ('3 routing cases', AMBER),
    ]
    for i, (label, c) in enumerate(callouts):
        x = 0.45 + i * 3.2
        add_rect(s, x, 3.2, 2.9, 0.9, fill=c)
        add_textbox(s, x, 3.3, 2.9, 0.5, label, size=14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER)
    add_textbox(s, 0.45, 4.6, 12, 0.35, 'Johnny, Jeng-lin Li · HKUST(GZ)',
                size=14, color=INK)
    add_textbox(s, 0.45, 5.0, 12, 0.3, 'Sept 2026 · ≈ 1,030 words + 1 figure + 6 references',
                size=11, color=MUTE)
    add_textbox(s, 0.45, 6.75, 11, 0.45,
                'For the SAAS design package: see /saas/index.html  ·  Flask prototype: http://127.0.0.1:5055',
                size=10, color=RGBColor(0xFF, 0xFF, 0xFF))

    # ----------------- 2 Problem
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '1 · Problem',
                  'Today, asking for two days off is an email exercise')
    add_bullets(s, 0.55, 1.45, 6.2, 5.2, [
        '• A student emails every instructor separately, attaches the same certificate several times, '
          'and negotiates each quiz / lab individually.',
        '• Each instructor improvises a reply — the programme office sees nothing; '
          'when a dispute arises, the Registry has no record of what was agreed.',
        '• Students face unpredictable outcomes and repeat private disclosures to people they barely know.',
        '• Instructors decide without context — they cannot see that the same student has already '
          'missed four sessions this term.',
        '• Programme offices and the Registry have no visibility, no audit trail, and no way to '
          'notice a struggling student before the damage reaches their grades.',
    ], size=13)

    add_rect(s, 6.95, 1.45, 6.0, 1.3, fill=RGBColor(0xFD, 0xF5, 0xE3))
    add_textbox(s, 7.05, 1.5, 5.8, 0.35, 'Why current tools do not help',
                size=12, bold=True, color=AMBER)
    add_bullets(s, 7.05, 1.85, 5.8, 0.9, [
        '• Form builders know nothing about timetables.',
        '• LMS attendance modules cannot distribute accommodation decisions.',
        '• Email creates no state anyone can audit.',
    ], size=11)

    add_rect(s, 6.95, 2.95, 6.0, 3.6, fill=RGBColor(0xE8, 0xF1, 0xFB))
    add_textbox(s, 7.05, 3.05, 5.8, 0.35, 'An absence decision is an academic decision',
                size=12, bold=True, color=BRAND)
    add_bullets(s, 7.05, 3.4, 5.8, 3.2, [
        '• Unmanaged conflict → missed assessment, disputed grade, or quiet disengagement.',
        '• The problem hits three groups at once: students, instructors, programme offices.',
        '• It is a genuine, verifiable problem on my own programme and at most universities.',
    ], size=11)
    footer(s, 2, 10)

    # ----------------- 3 Innovation
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '2 · Proposed Innovation',
                  'One submission replaces the whole email ritual')
    add_bullets(s, 0.55, 1.4, 6.2, 5.4, [
        '• Student selects category (medical, official activity, conference, family emergency, '
          'visa, bereavement, personal…), enters dates and reason, uploads evidence.',
        '• System resolves enrolments × assessment calendar → one conflict record per affected '
          'quiz / midterm / final / workshop / laboratory.',
        '• Routes by policy: A medical → instructor; B endorsed activity → roster auto-verified, '
          'instructors notified only; C > 7 days → programme office → Registry → instructors.',
        '• Instructor selects one of 5 accommodations per conflict: excused, make-up, alternative, '
          'weight redistribution, rejection with reasons.',
        '• Student sees outcome, deadline and make-up date in one timeline; notifications via email, '
          'Teams and push; every step audited.',
    ], size=13)

    add_rect(s, 6.95, 1.4, 6.0, 5.4, fill=RGBColor(0xE7, 0xF5, 0xEC))
    add_textbox(s, 7.05, 1.5, 5.8, 0.35, 'What is new (not the technologies)',
                size=12, bold=True, color=GREEN)
    add_bullets(s, 7.05, 1.85, 5.8, 4.8, [
        '• The technologies are mature — cloud hosting, campus SSO, workflow engine, '
          'document storage, integrations with SIS / LMS / email / Teams.',
        '• The new thing is the information design.',
        '• Unit of work = consequence, not request. Every affected assessment is resolved '
          'explicitly rather than left to negotiation.',
        '• Policy parameters live in Registry-owned configuration — rules change without a release.',
        '• Privacy is structural: instructors receive the accommodation to apply, '
          'never the medical reason or certificate behind it.',
    ], size=11)
    footer(s, 3, 10)

    # ----------------- 4 How it works (workflow fig)
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '3 · How It Works', 'One request → resolved conflicts → policy routing → accommodation')
    s.shapes.add_picture(FIG1, Inches(2.15), Inches(1.45), width=Inches(9.0))
    footer(s, 4, 10)

    # ----------------- 5 Routing cases
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '4 · Three Routing Cases', 'Configuration, not hard-coded logic')
    cases = [
        ('CASE A · Medical Leave', '< 7 days', '→ affected instructors decide',
         'Each instructor picks an accommodation; if >3 days Registry countersigns.', GREEN),
        ('CASE B · Official Activity', 'roster match', '→ auto-verified, instructors notified only',
         'No approval gate; students on the endorsed roster get the highest priority lane.', AMBER),
        ('CASE C · Absence > 7 days', 'long absence', '→ programme office → Registry → instructors',
         'Registry may attach conditions (coursework plan, minimum attendance) before final approval.', RED),
    ]
    for i, (name, code, route, detail, c) in enumerate(cases):
        x = 0.55 + i * 4.2
        add_rect(s, x, 1.5, 3.9, 0.7, fill=c)
        add_textbox(s, x, 1.55, 3.9, 0.5, name, size=14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER)
        add_rect(s, x, 2.25, 3.9, 3.2, fill=RGBColor(0xFA, 0xFA, 0xFA), line=LINE)
        add_textbox(s, x + 0.2, 2.35, 3.5, 0.3, 'Trigger', size=10, bold=True, color=MUTE)
        add_textbox(s, x + 0.2, 2.62, 3.5, 0.5, code, size=14, bold=True, color=INK)
        add_textbox(s, x + 0.2, 3.18, 3.5, 0.3, 'Route', size=10, bold=True, color=MUTE)
        add_textbox(s, x + 0.2, 3.45, 3.5, 0.7, route, size=12, color=INK)
        add_textbox(s, x + 0.2, 4.25, 3.5, 1.0, detail, size=10.5, color=MUTE, font='Calibri')
    add_textbox(s, 0.55, 5.8, 12.4, 0.4,
             'Case B students: auto-verified, no document required, highest priority lane · '
             'Case C must pass Registry before instructors act',
             size=11, color=MUTE)
    footer(s, 5, 10)

    # ----------------- 6 Course concept 1: complements
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '5 · Course Concept · Complements',
                  'Teece (1986): value is co-produced with the assets an innovator does not own')
    s.shapes.add_picture(FIG3, Inches(0.5), Inches(1.35), width=Inches(7.6))
    add_bullets(s, 8.3, 1.4, 4.6, 5.4, [
        '• SAAS value is almost entirely co-produced:',
        '   – without assessment calendar → conflict detection is guesswork',
        '   – without campus SSO → adoption collapses',
        '   – without email / Teams → decisions stall in a queue nobody opens',
        '   – without Registry-owned configuration → the system fossilises',
        '• Hard work is securing the complements, not writing the application.',
        '• Implication: prioritise integration contracts (SIS, LMS, IT) and '
          'policy ownership over feature scope in the first release.',
    ], size=11.5)
    footer(s, 6, 10)

    # ----------------- 7 Course concept 2: S-curve
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '6 · Course Concept · S-curve',
                  'Why SAAS is not "a better form" — Utterback & Abernathy (1975); Foster (1986)')
    s.shapes.add_picture(FIG2, Inches(0.5), Inches(1.35), width=Inches(8.0))
    add_bullets(s, 8.7, 1.4, 4.3, 5.4, [
        '• Rule-based workflow automation sits on the flat, mature part of its curve — '
          'competing on form features yields diminishing returns.',
        '• SAAS matters only by starting a new curve: joining academic-calendar data to a '
          'privacy-aware approval workflow.',
        '• Same logic explains the deliberate omission: AI triage and certificate-authenticity '
          'checking sit on the steep part of the LLM curve.',
        '• A wrong automated decision in a high-stakes academic process is unacceptable; '
          'these features wait until reliability crosses the threshold.',
    ], size=11)
    footer(s, 7, 10)

    # ----------------- 8 Business value
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '7 · Business Value', 'An institutional licence; measurable cost becomes visible')
    # left bullets
    add_bullets(s, 0.55, 1.4, 7.6, 5.4, [
        '• Customer & payer: the university, not the student — Registry and IT office '
          'buy an annual licence sized by student population.',
        '• Users (students, instructors) pay in time, not money.',
        '• Value: makes an unmeasured cost measurable.',
        '   – Instructor hours on absence email',
        '   – Administrative effort on disputes',
        '   – Compliance risk of holding no record',
        '• Differentiator: policy-configuration layer + portability — same product serves '
          'HKUST(GZ), HKUST, CUHK(SZ) and other institutions whose thresholds differ.',
        '• Needed complements: campus SSO, SIS / LMS vendors, Microsoft for delivery, '
          'cross-campus policy owners willing to harmonise.',
    ], size=12)

    add_rect(s, 8.35, 1.4, 4.6, 5.4, fill=RGBColor(0xF0, 0xEB, 0xFA))
    add_textbox(s, 8.45, 1.5, 4.4, 0.35, 'Why choose it over alternatives?',
                size=12, bold=True, color=PURPLE)
    add_bullets(s, 8.45, 1.85, 4.4, 4.7, [
        '• Generic form tool → cannot see the timetable.',
        '• LMS attendance module → records absence, but cannot distribute an accommodation decision.',
        '• Building in-house → hard-codes one university\'s policy.',
        '• SAAS → auditable decision history + configurable policy + portable across institutions.',
    ], size=11)
    footer(s, 8, 10)

    # ----------------- 9 Social impact + risks
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '8 · Social Impact & Risks', 'Equity gains are uneven; data is a double-edged good')
    # positive
    add_rect(s, 0.55, 1.4, 6.0, 4.6, fill=RGBColor(0xE7, 0xF5, 0xEC))
    add_textbox(s, 0.65, 1.5, 5.8, 0.35, 'Positive impact', size=12, bold=True, color=GREEN)
    add_bullets(s, 0.65, 1.85, 5.8, 4.0, [
        '• Equity: outcome stops depending on which instructor you ask or how persuasive '
          'your email is.',
        '• Recurring-need students (chronic illness, caregiving, visa, athletes) gain most — '
          'the system carries context instead of forcing retellings.',
        '• Early-warning rules can reach a struggling student before failure appears in grades.',
        '• Unavoidable date absences (visa appointments) cannot be rejected on grounds of date.',
    ], size=11)

    # risks
    add_rect(s, 6.75, 1.4, 6.2, 4.6, fill=RGBColor(0xFB, 0xE9, 0xEA))
    add_textbox(s, 6.85, 1.5, 6.0, 0.35, 'Risks & mitigations', size=12, bold=True, color=RED)
    add_bullets(s, 6.85, 1.85, 6.0, 4.0, [
        '• Gains are uneven: well-connected students adopt earliest; '
          'need an assisted non-digital route.',
        '• Surveillance: risk scoring may stigmatise; medical documents are among the most '
          'sensitive records a university holds.',
        '• Automation bias: instructors may ratify whatever the queue presents.',
        '• Employment effect: administrative staff shift from email triage to case management '
          '(a change in skills, not simply less work).',
        '• Mitigations built in: instructors never see medical reasons; every decision stays '
          'human; retention periods fixed; AI is advisory, logged and refusable.',
    ], size=11)
    footer(s, 9, 10)

    # ----------------- 10 Reflection + end
    s = prs.slides.add_slide(BLANK)
    add_title_bar(s, '9 · Final Reflection', 'Worth developing — but narrowed')
    add_bullets(s, 0.55, 1.4, 12.4, 3.5, [
        '• Yes, develop further as a team project — the problem is real and bounded, the '
          'complements already exist on campus, and its hardest questions are organisational '
          'rather than technical, which makes it a genuine test of collaborative innovation.',
        '• But narrow scope: the full system is too large for one term.',
        '• Recommend scoping the team project to the assessment-conflict workflow, where both '
          'the value and the privacy tension are sharpest.',
    ], size=14)

    add_rect(s, 0.55, 5.2, 12.4, 1.5, fill=RGBColor(0xF5, 0xF5, 0xF5))
    add_textbox(s, 0.7, 5.3, 12, 0.35, 'References', size=12, bold=True, color=INK)
    refs = ('Foster (1986); Teece (1986); Utterback & Abernathy (1975); '
            'Microsoft Entra ID docs (2025); GDPR (2016); FERPA (US Dept. of Education).')
    add_textbox(s, 0.7, 5.65, 12, 1.0, refs, size=11, color=MUTE)
    footer(s, 10, 10)

    prs.save(PPTX)
    # attach speaker notes by slide number
    for i, sl in enumerate(prs.slides, 1):
        if i in NOTES:
            sl.notes_slide.notes_text_frame.text = NOTES[i]
    prs.save(PPTX)
    print('saved', PPTX)


if __name__ == '__main__':
    build()